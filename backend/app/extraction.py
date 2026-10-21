"""Bounded in-memory extraction. No original filenames are used as filesystem paths."""

import csv
import io
import json
import re
import warnings
import zipfile
from pathlib import Path
from defusedxml import ElementTree as SafeET
from pydantic import ValidationError
from .models import ExtractionInfo, RawObservation
from .parsing import ParseAccumulator, deduplicate, parse_rows, parse_text

MAX_FILE = 4_000_000
MAX_PAGES = 15
MAX_TEXT = 150_000
MAX_ZIP_EXPANDED = 16_000_000
ALLOWED = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".txt",
    ".csv",
    ".tsv",
    ".json",
    ".xlsx",
    ".docx",
}


class ExtractionError(ValueError):
    def __init__(self, message, code="invalid_report", status=422):
        super().__init__(message)
        self.code = code
        self.status = status


def safe_zip(data):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ExtractionError("This is not a valid Office document.")
    members = z.infolist()
    if len(members) > 1000 or sum(x.file_size for x in members) > MAX_ZIP_EXPANDED:
        z.close()
        raise ExtractionError(
            "Document expansion exceeds the safe processing limit. Export a smaller PDF or CSV."
        )
    if any(
        x.flag_bits & 1 or x.filename.startswith("/") or ".." in Path(x.filename).parts
        for x in members
    ):
        z.close()
        raise ExtractionError(
            "Encrypted or unsafe document container is not supported."
        )
    if any("vbaproject" in x.filename.casefold() for x in members):
        z.close()
        raise ExtractionError(
            "Macro-enabled documents are not supported. Export a PDF or CSV."
        )
    return z


def decode_text(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8-sig"
    try:
        text = data.decode(encoding)
    except UnicodeError:
        raise ExtractionError(
            "Text must be UTF-8 or UTF-16. Re-export the report as UTF-8 or PDF."
        )
    if "\x00" in text:
        raise ExtractionError("This file is not a supported text document.")
    if len(text) > MAX_TEXT:
        raise ExtractionError("Report text is too long. Split it into smaller files.")
    return text


def multiple_patient_ids(text):
    found = re.findall(
        r"(?im)^\s*patient\s*(?:id|identifier|no\.?)\s*[:#]\s*([^\n]{1,60})", text
    )
    return len({x.strip().casefold() for x in found}) > 1


def extract(data: bytes, filename: str, allow_cloud=False):
    if not data:
        raise ExtractionError("The uploaded file is empty.")
    if len(data) > MAX_FILE:
        raise ExtractionError("Files must be 4 MB or smaller.", status=413)
    filename = Path(filename.replace("\\", "/")).name[:120]
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ExtractionError(
            "Unsupported file. Use PDF, PNG, JPG, WebP, TXT, CSV, TSV, JSON, XLSX, or DOCX.",
            status=415,
        )
    acc = ParseAccumulator()
    info = ExtractionInfo(filename=filename, format=suffix[1:])
    full_text = ""
    needs_ocr = False
    if suffix == ".pdf":
        if b"%PDF-" not in data[:1024]:
            raise ExtractionError("The file does not have a valid PDF signature.")
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ExtractionError(
                    "Password-protected PDFs are not supported. Upload an unlocked copy."
                )
            count = len(reader.pages)
            if not 1 <= count <= MAX_PAGES:
                raise ExtractionError(
                    f"PDFs must contain 1–{MAX_PAGES} pages. Split this report before uploading."
                )
            info.page_count = count
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError(
                "The PDF is damaged or could not be opened. Re-export it and try again."
            )
        import pdfplumber

        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for pageno, page in enumerate(pdf.pages, 1):
                    # Limit pathological page geometry before layout/table analysis.
                    if (
                        page.width <= 0
                        or page.height <= 0
                        or page.width > 5000
                        or page.height > 5000
                    ):
                        raise ExtractionError(
                            "PDF page dimensions exceed safe processing limits."
                        )
                    try:
                        # Coordinates and table cells are complementary, not competing guesses.
                        tables = page.extract_tables(
                            table_settings={
                                "vertical_strategy": "lines",
                                "horizontal_strategy": "lines",
                            }
                        )
                        if not tables:
                            tables = page.extract_tables(
                                table_settings={
                                    "vertical_strategy": "text",
                                    "horizontal_strategy": "text",
                                    "min_words_vertical": 2,
                                    "min_words_horizontal": 1,
                                }
                            )
                        for tableindex, table in enumerate(tables, 1):
                            parse_rows(
                                table, acc, "table", pageno, f"Table {tableindex}"
                            )
                    except Exception:
                        acc.warnings.append(
                            f"Page {pageno}: table extraction was incomplete; text extraction was attempted."
                        )
                    text = (
                        page.extract_text(layout=False, x_tolerance=2, y_tolerance=3)
                        or ""
                    )
                    if not text.strip():
                        try:
                            text = (
                                reader.pages[pageno - 1].extract_text(
                                    extraction_mode="layout"
                                )
                                or ""
                            )
                        except Exception:
                            text = ""
                    full_text += text + "\n"
                    if len(full_text) > MAX_TEXT:
                        raise ExtractionError(
                            "The extracted PDF text exceeds the processing limit. Split the report."
                        )
                    parse_text(text, acc, pageno)
                    if len(text.strip()) < 30:
                        needs_ocr = True
                        acc.warnings.append(
                            f"Page {pageno} has little or no readable text. Its results may be missing without OCR."
                        )
            info.methods = [
                "PDF tables",
                "PDF text + coordinate layout",
                "PDF text fallback",
            ]
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError(
                "PDF extraction failed. Try a fresh PDF export or a clear image."
            )
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image

        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                image = Image.open(io.BytesIO(data))
                expected = {
                    ".png": "PNG",
                    ".jpg": "JPEG",
                    ".jpeg": "JPEG",
                    ".webp": "WEBP",
                }[suffix]
                if image.format != expected:
                    raise ExtractionError(
                        "Image content does not match the selected file extension."
                    )
                if image.width * image.height > 16_000_000:
                    raise ExtractionError("Images must be 16 megapixels or smaller.")
                if getattr(image, "n_frames", 1) > 1:
                    raise ExtractionError(
                        "Animated or multi-frame images are not supported."
                    )
                image.verify()
            # Re-encode before provider processing: remove EXIF and cap dimensions.
            image = Image.open(io.BytesIO(data)).convert("RGB")
            image.thumbnail((2400, 2400))
            out = io.BytesIO()
            image.save(out, format="JPEG", quality=90)
            data = out.getvalue()
            suffix = ".jpg"
            needs_ocr = True
            info.page_count = 1
            info.methods = ["Image validation"]
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError(
                "The image is damaged, too large, or not a supported image."
            )
    elif suffix in {".csv", ".tsv"}:
        full_text = decode_text(data)
        try:
            dialect = (
                csv.excel_tab
                if suffix == ".tsv"
                else csv.Sniffer().sniff(full_text[:8192], delimiters=",;\t")
            )
        except csv.Error:
            dialect = csv.excel
        try:
            parse_rows(csv.reader(io.StringIO(full_text), dialect), acc, "csv")
        except csv.Error:
            raise ExtractionError("The CSV/TSV structure is invalid.")
        info.methods = ["Delimited columns"]
    elif suffix == ".txt":
        full_text = decode_text(data)
        parse_text(full_text, acc)
        info.methods = ["Anchored text parsing"]
    elif suffix == ".json":
        full_text = decode_text(data)
        try:
            obj = json.loads(full_text)
            if isinstance(obj, dict) and obj.get("resourceType") in {
                "Bundle",
                "Observation",
                "DiagnosticReport",
            }:
                from .fhir import extract_fhir

                acc.items, fhir_warnings = extract_fhir(obj)
                acc.warnings.extend(fhir_warnings)
                info.methods = ["FHIR observations / components"]
            else:
                rows = obj.get("observations") if isinstance(obj, dict) else obj
                if not isinstance(rows, list) or len(rows) > 300:
                    raise ValueError()
                for i, row in enumerate(rows, 1):
                    if not isinstance(row, dict):
                        raise ValueError()
                    row = dict(row)
                    row["source_kind"] = "json"
                    row["source_row"] = i
                    row["user_verified"] = False
                    row.setdefault(
                        "source_text",
                        json.dumps(
                            {
                                k: v
                                for k, v in row.items()
                                if k not in {"source_text", "issues"}
                            },
                            ensure_ascii=False,
                        )[:600],
                    )
                    acc.items.append(RawObservation.model_validate(row))
                info.methods = ["Validated observation JSON"]
        except (
            ValueError,
            TypeError,
            AttributeError,
            KeyError,
            IndexError,
            ValidationError,
            RecursionError,
        ):
            raise ExtractionError(
                "JSON must contain supported observation objects or a FHIR Observation/Bundle. See the documented schema."
            )
    elif suffix == ".xlsx":
        with safe_zip(data) as z:
            if "xl/workbook.xml" not in z.namelist():
                raise ExtractionError("This is not an XLSX workbook.")
        from openpyxl import load_workbook

        try:
            book = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            if len(book.worksheets) > 8:
                raise ExtractionError("Workbooks must have 8 sheets or fewer.")
            for sheet in book.worksheets:
                if (
                    sheet.max_row
                    and sheet.max_row > 5000
                    or sheet.max_column
                    and sheet.max_column > 40
                ):
                    raise ExtractionError(
                        "Workbook dimensions exceed the processing limit."
                    )
                rows = []
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if i > 5000:
                        raise ExtractionError("Workbook row limit reached.")
                    rows.append([str(x) if x is not None else "" for x in row[:40]])
                parse_rows(rows, acc, "xlsx", section=sheet.title[:100])
            book.close()
            info.methods = ["Read-only workbook columns"]
            acc.warnings.append(
                "Spreadsheet formulas are not calculated. Only stored cell values are read; verify computed or date-formatted cells."
            )
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError(
                "The workbook could not be read. Export its results as CSV or PDF."
            )
    elif suffix == ".docx":
        with safe_zip(data) as z:
            try:
                xml = z.read("word/document.xml")
            except KeyError:
                raise ExtractionError("This is not a DOCX document.")
        try:
            root = SafeET.fromstring(xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for tableindex, table in enumerate(root.findall(".//w:tbl", ns), 1):
                rows = []
                for row in table.findall("w:tr", ns):
                    rows.append(
                        [
                            " ".join(t.text or "" for t in cell.findall(".//w:t", ns))
                            for cell in row.findall("w:tc", ns)
                        ]
                    )
                parse_rows(rows, acc, "docx", section=f"Table {tableindex}")
            full_text = "\n".join(
                "".join(t.text or "" for t in p.findall(".//w:t", ns))
                for p in root.findall(".//w:p", ns)
            )
            if len(full_text) > MAX_TEXT:
                raise ExtractionError("Document text exceeds the processing limit.")
            paragraphs = "\n".join(
                "".join(t.text or "" for t in p.findall(".//w:t", ns))
                for p in root.findall("w:body/w:p", ns)
            )
            parse_text(paragraphs, acc, kind="docx")
            info.methods = ["DOCX tables", "DOCX paragraphs"]
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError("The DOCX document could not be read.")
    if needs_ocr:
        if allow_cloud:
            if (info.page_count or 1) > 8:
                raise ExtractionError(
                    "Cloud OCR supports up to 8 pages per request. Split this scanned report."
                )
            from .cloud import CloudUnavailable, extract_with_cloud

            try:
                ocr = extract_with_cloud(
                    data, "application/pdf" if suffix == ".pdf" else "image/jpeg"
                )
                acc.items.extend(ocr)
                info.cloud_used = True
                info.methods.append("Google Gemini OCR (unverified)")
                acc.warnings.append(
                    "Cloud OCR is probabilistic. Every cloud-extracted observation requires manual verification before interpretation."
                )
            except CloudUnavailable as exc:
                if not acc.items:
                    raise
                acc.warnings.append(
                    "Cloud OCR failed. Only readable digital content was retained; scanned pages remain incomplete. "
                    + str(exc)
                )
        else:
            acc.warnings.append(
                "Scanned/image content was not sent to a third party. Enable cloud processing only if you consent, or paste/export the report text."
            )
    if multiple_patient_ids(full_text):
        acc.warnings.append(
            "Multiple patient identifiers were detected. This may contain reports for different people; separate and verify the records."
        )
        for item in acc.items:
            item.issues.append(
                "Multiple patient identifiers detected; verify the observation belongs to the intended person."
            )
    info.warnings = list(dict.fromkeys(acc.warnings))[:100]
    info.ignored_line_count = acc.ignored
    info.unparsed_lines = acc.unparsed
    return deduplicate(acc.items)[:300], info
